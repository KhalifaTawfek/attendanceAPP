import io
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook

from tests.conftest import auth_headers, png_bytes


def iso(dt: datetime) -> str:
    return dt.isoformat()


def make_event(client, admin, start_offset_min=-5, duration_min=60, **extra):
    now = datetime.now(UTC)
    body = {
        "title": "Standup",
        "location": "Room 1",
        "starts_at": iso(now + timedelta(minutes=start_offset_min)),
        "ends_at": iso(now + timedelta(minutes=start_offset_min + duration_min)),
        **extra,
    }
    r = client.post("/api/events", headers=admin, json=body)
    assert r.status_code == 201, r.text
    return r.json()[0]


def submit(client, headers, event_id, declaration="true", photo=None):
    return client.post(
        "/api/attendances",
        headers=headers,
        data={"event_id": event_id, "activity_summary": "Attended the standup",
              "declaration": declaration, "location_note": "Desk 4"},
        files={"photo": ("me.png", photo if photo is not None else png_bytes(), "image/png")},
    )


def test_health_and_metrics(client):
    assert client.get("/health").json()["status"] == "ok"
    r = client.get("/metrics")
    assert r.status_code == 200 and "http_requests_total" in r.text


def test_auth_required_and_admin_only(client, user):
    assert client.get("/api/auth/me").status_code == 401
    assert client.get("/api/admin/users", headers=user).status_code == 403
    assert client.get("/api/auth/me", headers=user).json()["email"] == "jane@test.io"


def test_login_bad_password(client):
    r = client.post("/api/auth/login", json={"email": "admin@test.io", "password": "nope"})
    assert r.status_code == 401


def test_self_registration_disabled_by_default(client):
    r = client.post("/api/auth/register", json={"email": "x@test.io", "full_name": "X", "password": "12345678"})
    assert r.status_code == 403


def test_event_validation(client, admin):
    now = datetime.now(UTC)
    r = client.post("/api/events", headers=admin, json={
        "title": "Bad", "starts_at": iso(now), "ends_at": iso(now - timedelta(hours=1))})
    assert r.status_code == 422
    r = client.post("/api/events", headers=admin, json={
        "title": "Naive", "starts_at": "2026-01-01T10:00:00", "ends_at": "2026-01-01T11:00:00"})
    assert r.status_code == 422


def test_event_defaults(client, admin):
    ev = make_event(client, admin)
    opens = datetime.fromisoformat(ev["check_in_opens_at"])
    starts = datetime.fromisoformat(ev["starts_at"])
    assert starts - opens == timedelta(minutes=15)
    assert ev["late_after_minutes"] == 10


def test_full_attendance_flow(client, admin, user):
    ev = make_event(client, admin)
    events = client.get("/api/events", headers=user).json()
    assert events[0]["check_in_open"] is True and events[0]["my_attendance_status"] is None

    r = submit(client, user, ev["id"])
    assert r.status_code == 201, r.text
    att = r.json()
    assert att["status"] == "pending" and att["is_late"] is False

    # duplicate rejected
    assert submit(client, user, ev["id"]).status_code == 409

    # own photo visible, re-encoded as JPEG
    r = client.get(f"/api/attendances/{att['id']}/photo", headers=user)
    assert r.status_code == 200 and r.content[:2] == b"\xff\xd8"

    # admin review
    r = client.patch(f"/api/admin/attendances/{att['id']}", headers=admin,
                     json={"status": "approved", "review_comment": "ok"})
    assert r.json()["status"] == "approved"
    assert client.get("/api/attendances/me", headers=user).json()[0]["status"] == "approved"

    stats = client.get("/api/admin/stats", headers=admin).json()
    assert stats["approved"] == 1 and stats["total_users"] == 1


def test_late_flag(client, admin, user):
    ev = make_event(client, admin, start_offset_min=-30)
    assert submit(client, user, ev["id"]).json()["is_late"] is True


def test_check_in_closed(client, admin, user):
    future = make_event(client, admin, start_offset_min=120)
    assert submit(client, user, future["id"]).status_code == 409


def test_declaration_required_and_bad_image(client, admin, user):
    ev = make_event(client, admin)
    assert submit(client, user, ev["id"], declaration="false").status_code == 422
    assert submit(client, user, ev["id"], photo=b"not an image").status_code == 422


def test_photo_privacy(client, admin, user):
    ev = make_event(client, admin)
    att = submit(client, user, ev["id"]).json()
    client.post("/api/admin/users", headers=admin,
                json={"email": "bob@test.io", "full_name": "Bob", "password": "bobpass123"})
    bob = auth_headers(client, "bob@test.io", "bobpass123")
    assert client.get(f"/api/attendances/{att['id']}/photo", headers=bob).status_code == 404
    assert client.get(f"/api/attendances/{att['id']}/photo", headers=admin).status_code == 200


def test_filters_absentees_and_export(client, admin, user):
    ev = make_event(client, admin)
    submit(client, user, ev["id"])
    client.post("/api/admin/users", headers=admin,
                json={"email": "bob@test.io", "full_name": "Bob", "password": "bobpass123"})

    page = client.get("/api/admin/attendances", headers=admin, params={"q": "jane", "status": "pending"}).json()
    assert page["total"] == 1
    assert client.get("/api/admin/attendances", headers=admin, params={"status": "approved"}).json()["total"] == 0

    absent = client.get(f"/api/admin/events/{ev['id']}/absentees", headers=admin).json()
    assert [u["email"] for u in absent] == ["bob@test.io"]

    csv_r = client.get("/api/admin/attendances/export", headers=admin, params={"format": "csv"})
    assert csv_r.status_code == 200 and "Jane Doe" in csv_r.text
    assert "attachment" in csv_r.headers["content-disposition"]

    x = client.get("/api/admin/attendances/export", headers=admin, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.max_row == 2 and ws["D2"].value == "Jane Doe"


def test_deactivated_user_cannot_login(client, admin, user):
    uid = client.get("/api/auth/me", headers=user).json()["id"]
    client.patch(f"/api/admin/users/{uid}", headers=admin, json={"is_active": False})
    assert client.get("/api/auth/me", headers=user).status_code == 401
    r = client.post("/api/auth/login", json={"email": "jane@test.io", "password": "janepass123"})
    assert r.status_code == 403


def test_admin_cannot_deactivate_self(client, admin):
    me = client.get("/api/auth/me", headers=admin).json()
    assert client.patch(f"/api/admin/users/{me['id']}", headers=admin, json={"is_active": False}).status_code == 400
